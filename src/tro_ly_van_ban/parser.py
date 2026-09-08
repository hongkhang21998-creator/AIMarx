from io import BytesIO
from zipfile import ZipFile
from docx import Document
from pypdf import PdfReader

MAX_BYTES = 10 * 1024 * 1024
MAX_TEXT = 60000


def parse(data: bytes, suffix: str) -> tuple[list[dict], list[str]]:
    warnings = []
    if suffix == ".txt":
        items = [("Đoạn", x) for x in data.decode("utf-8-sig").splitlines() if x.strip()]
    elif suffix == ".pdf":
        pdf = PdfReader(BytesIO(data))
        if pdf.is_encrypted:
            raise ValueError("PDF có mật khẩu chưa được hỗ trợ")
        if len(pdf.pages) > 100:
            raise ValueError("Giới hạn 100 trang PDF")
        items = []
        for i, page in enumerate(pdf.pages, 1):
            value = page.extract_text() or ""
            items.append((f"Trang {i}", value))
            if not value.strip():
                warnings.append(f"Trang {i}: cần OCR; chưa đọc được chữ")
    elif suffix == ".docx":
        with ZipFile(BytesIO(data)) as archive:
            if sum(x.file_size for x in archive.infolist()) > 40 * 1024 * 1024:
                raise ValueError("DOCX giải nén vượt 40 MB")
        doc = Document(BytesIO(data))
        items = [(f"Đoạn {i}", p.text) for i, p in enumerate(doc.paragraphs, 1) if p.text.strip()]
        for i, table in enumerate(doc.tables, 1):
            for j, row in enumerate(table.rows, 1):
                items.append((f"Bảng {i}, hàng {j}", " | ".join(c.text for c in row.cells)))
        warnings.append("DOCX: chưa đọc textbox, ảnh, header/footer hoặc chú thích")
    else:
        raise ValueError("Chỉ nhận PDF, DOCX và TXT UTF-8")
    if sum(len(t) for _, t in items) > MAX_TEXT:
        raise ValueError("Nội dung vượt 60.000 ký tự; hãy chia tài liệu")
    blocks = [{"id": f"b{i}", "location": label, "text": text} for i, (label, text) in enumerate(items, 1) if text.strip()]
    return blocks, warnings
