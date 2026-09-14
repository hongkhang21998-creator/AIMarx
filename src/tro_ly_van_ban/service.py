import asyncio
import contextlib
import hashlib
import json
import os
import sqlite3
import threading
from pathlib import Path
from typing import TypedDict
from docx import Document
from langgraph.graph import StateGraph, START, END
from .domain import Extraction, validate_evidence
from .fsguard import freeze, harden_dir, thaw
from .model import extract, ModelUnavailable
from .parser import parse, MAX_BYTES
from .storage_guard import check_data_root
from .token_usage import TokenUsage, capture_usage


class NotFound(ValueError):
    """Không có tài nguyên được yêu cầu — HTTP 404."""


class Conflict(ValueError):
    """Trạng thái đã đổi dưới chân người gửi — HTTP 409.

    Khác với lỗi nhập liệu: yêu cầu đúng cú pháp, chỉ là nó nói về một phiên bản
    không còn là hiện hành. Gửi lại y nguyên vẫn hỏng; phải tải lại trang trước.
    """


def requires_ocr(blocks: list[dict], warnings: list[str]) -> bool:
    """Điều kiện cần OCR, suy ra từ dữ liệu bất biến của lần nhập.

    Cố ý không đọc cột `state`: state bị ghi đè bởi save/run, nên dùng nó làm
    hàng rào thì chỉ cần một lần ghi là cửa mở vĩnh viễn. blocks và warnings
    được chốt lúc ingest và không đổi.
    """
    return not blocks or any("cần OCR" in w for w in warnings)


class Flow(TypedDict):
    blocks: list[dict]
    mode: str
    model: str
    result: dict


def infer(state: Flow):
    return {"result": extract(state["blocks"], state["mode"], state["model"]).model_dump()}


def validate(state: Flow):
    validate_evidence(Extraction.model_validate(state["result"]), state["blocks"])
    return {}


builder = StateGraph(Flow)
builder.add_node("extract", infer)
builder.add_node("validate", validate)
builder.add_edge(START, "extract")
builder.add_edge("extract", "validate")
builder.add_edge("validate", END)
graph = builder.compile()


class Service:
    def __init__(self, root: str | Path, mode="ollama", model="qwen3:0.6b", required_mount=None):
        if mode not in {"ollama", "demo"}:
            raise ValueError("TLVB_MODE must be ollama or demo")
        # Phai kiem truoc mkdir: o vang mat thi mkdir(parents=True) se dung kho rong tren tmpfs.
        check_data_root(root, required_mount)
        self.root = Path(root).absolute()
        if self.root.is_symlink():
            raise ValueError("Data directory cannot be a symlink")
        self.root.mkdir(parents=True, exist_ok=True, mode=0o700)
        self.root = self.root.resolve()
        # Windows bo qua mode cua mkdir; canh bao thay vi mat quyen rieng tu trong im lang.
        self.warnings = [w for w in [harden_dir(self.root)] if w]
        for folder in ("originals", "drafts"):
            path = self.root / folder
            if path.is_symlink():
                raise ValueError("Storage directory cannot be a symlink")
            path.mkdir(exist_ok=True, mode=0o700)
            harden_dir(path)
        if (self.root / "state.sqlite3").is_symlink():
            raise ValueError("Database cannot be a symlink")
        self.mode, self.model = mode, model
        self.lock = threading.RLock()
        with self.db() as db:
            db.executescript('''
            CREATE TABLE IF NOT EXISTS documents(id TEXT PRIMARY KEY, name TEXT, suffix TEXT, blocks TEXT, warnings TEXT, state TEXT, error TEXT, error_kind TEXT NOT NULL DEFAULT '');
            CREATE TABLE IF NOT EXISTS versions(document_id TEXT, version INTEGER, content TEXT, hash TEXT, mode TEXT, model TEXT, PRIMARY KEY(document_id,version));
            CREATE TABLE IF NOT EXISTS approvals(document_id TEXT, version INTEGER, hash TEXT, action TEXT, reason TEXT, created TEXT DEFAULT CURRENT_TIMESTAMP);
            ''')

            if "draft_hash" not in {row[1] for row in db.execute("PRAGMA table_info(versions)")}:
                db.execute("ALTER TABLE versions ADD COLUMN draft_hash TEXT NOT NULL DEFAULT ''")
            if "error_kind" not in {row[1] for row in db.execute("PRAGMA table_info(documents)")}:
                db.execute("ALTER TABLE documents ADD COLUMN error_kind TEXT NOT NULL DEFAULT ''")
            if "classification" not in {row[1] for row in db.execute("PRAGMA table_info(documents)")}:
                db.execute("ALTER TABLE documents ADD COLUMN classification TEXT NOT NULL DEFAULT 'unknown'")
        self.token_usage = TokenUsage(self.db)

    @contextlib.contextmanager
    def db(self):
        # Phai close: `with conn` chi commit/rollback. Windows giu file handle,
        # khoa DB va chan xoa tep neu connection khong dong.
        conn = sqlite3.connect(self.root / "state.sqlite3", timeout=30)
        conn.row_factory = sqlite3.Row
        try:
            with conn:
                yield conn
        finally:
            conn.close()

    def path(self, folder, name):
        path = self.root / folder / name
        if path.is_symlink() or path.resolve().parent != (self.root / folder).resolve():
            raise ValueError("Đường dẫn kho không hợp lệ")
        return path

    def listing(self):
        with self.db() as db:
            return [dict(x) for x in db.execute("SELECT id,name,state,error,error_kind FROM documents ORDER BY rowid DESC")]

    def get(self, document_id):
        with self.db() as db:
            row = db.execute("SELECT * FROM documents WHERE id=?", (document_id,)).fetchone()
            if row is None:
                raise NotFound("Không tìm thấy tài liệu")
            result = dict(row)
            result["blocks"] = json.loads(result["blocks"])
            result["warnings"] = json.loads(result["warnings"])
            latest = db.execute("SELECT * FROM versions WHERE document_id=? ORDER BY version DESC LIMIT 1", (document_id,)).fetchone()
            result["latest"] = dict(latest) if latest else None
            result["approvals"] = [dict(x) for x in db.execute("SELECT * FROM approvals WHERE document_id=?", (document_id,))]
            return result

    def ingest(self, name, data, classification="unknown"):
        if type(classification) is not str or classification not in {"unknown", "internal", "restricted", "public", "synthetic"}:
            raise ValueError("Phân loại tài liệu không hợp lệ")
        if not data or len(data) > MAX_BYTES:
            raise ValueError("File rỗng hoặc vượt 10 MB")
        suffix = Path(name).suffix.lower()
        if suffix not in {".txt", ".pdf", ".docx"}:
            raise ValueError("Chỉ nhận PDF, DOCX, TXT")
        digest = hashlib.sha256(data).hexdigest()
        with self.lock:
            with self.db() as db:
                if db.execute("SELECT 1 FROM documents WHERE id=?", (digest,)).fetchone():
                    return digest
            try:
                blocks, warnings = parse(data, suffix)
            except Exception as exc:
                raise ValueError("Không đọc được file: " + str(exc)[:200]) from exc
            target = self.path("originals", digest + suffix)
            if target.exists():
                if hashlib.sha256(target.read_bytes()).hexdigest() != digest:
                    raise ValueError("Bản gốc không khớp hash")
            else:
                with target.open("xb") as out:
                    out.write(data)
                freeze(target)
            state = "needs_ocr" if requires_ocr(blocks, warnings) else "ready"
            with self.db() as db:
                db.execute("INSERT INTO documents(id,name,suffix,blocks,warnings,state,error,error_kind) VALUES(?,?,?,?,?,?,?,?)", (digest, Path(name).name[:200], suffix, json.dumps(blocks, ensure_ascii=False), json.dumps(warnings, ensure_ascii=False), state, "", ""))
                db.execute("UPDATE documents SET classification=? WHERE id=?", (classification, digest))
        return digest

    def run(self, document_id, expected_version=None):
        """Trích xuất bằng model rồi lưu thành phiên bản mới.

        `expected_version` là phiên bản mà người gửi đang nhìn thấy. Kiểm nó
        **trước** khi gọi model: một tab mở từ hôm qua không được phép đẩy model
        chạy rồi đè lên bản người khác vừa sửa. Kiểm dưới cùng một lần giữ lock
        với lần ghi, nếu không thì giữa kiểm và ghi vẫn còn khe cho bản khác chen.
        """
        with self.lock:
            doc = self.get(document_id)
            if requires_ocr(doc["blocks"], doc["warnings"]):
                raise ValueError("Cần OCR đầy đủ trước khi xử lý")
            current = doc["latest"]["version"] if doc["latest"] else 0
            if expected_version is not None and expected_version != current:
                raise Conflict(f"Phiên bản đã thay đổi (bạn đang xem {expected_version}, hiện tại là {current}); tải lại trang")
            try:
                from fastmcp import Client
                from .mcp_server import create_mcp
                async def read_source():
                    async with Client(create_mcp(self)) as client:
                        response = await client.call_tool("read_document", {"document_id": document_id})
                        return response.data
                source = asyncio.run(read_source())
                with capture_usage(self.token_usage):
                    result = graph.invoke({"blocks": source["blocks"], "mode": self.mode, "model": self.model})
                # Chot lai lan nua o cua ghi. Lock dang giu nen khong the lech,
                # nhung the la hang rao khong phu thuoc vao viec ai do sau nay
                # van giu lock suot lan goi model.
                return self.save(document_id, result["result"], expected_version=current, provenance=self.mode)
            except Exception as exc:
                self.record_failure(document_id, exc)
                raise

    def record_failure(self, document_id, exc):
        """Ghi kết quả lần chạy vào error/error_kind, không đụng cột state.

        `state` mô tả phiên bản hiện hành: chờ duyệt, đã duyệt hay đã từ chối.
        Một lần trích xuất hỏng không tạo ra phiên bản mới và cũng không thu hồi
        xác nhận của phiên bản đang có, nên nó không được quyền ghi vào cột đó —
        ghi vào là mất chữ ký đã có mà không ai bấm nút thu hồi.

        Ngoại lệ là tài liệu chưa có phiên bản nào: ở đó không có trạng thái
        duyệt nào để giữ, lỗi chính là tình trạng hiện tại của tài liệu.
        """
        kind = "model_unavailable" if isinstance(exc, ModelUnavailable) else "error"
        message = str(exc)[:500]
        with self.lock, self.db() as db:
            if db.execute("SELECT 1 FROM versions WHERE document_id=? LIMIT 1", (document_id,)).fetchone():
                db.execute("UPDATE documents SET error=?,error_kind=? WHERE id=?", (message, kind, document_id))
            else:
                db.execute("UPDATE documents SET state=?,error=?,error_kind=? WHERE id=?", (kind, message, kind, document_id))

    def save(self, document_id, content, expected_version=None, provenance="manual"):
        with self.lock:
            doc = self.get(document_id)
            # Cua ghi duy nhat: run/manual/edit/save deu di qua day. Chan o giao dien
            # khong du vi cac route van nhan POST truc tiep.
            if requires_ocr(doc["blocks"], doc["warnings"]):
                raise ValueError("Cần OCR đầy đủ trước khi lập phiếu")
            current = doc["latest"]["version"] if doc["latest"] else 0
            if expected_version is not None and expected_version != current:
                raise Conflict(f"Phiên bản đã thay đổi (bạn đang xem {expected_version}, hiện tại là {current}); tải lại trang")
            value = Extraction.model_validate(content)
            validate_evidence(value, doc["blocks"])
            encoded = json.dumps(value.model_dump(), ensure_ascii=False, sort_keys=True)
            digest = hashlib.sha256(encoded.encode()).hexdigest()
            version = current + 1
            draft = Document()
            draft.add_heading("DỰ THẢO PHIẾU XỬ LÝ — CHỜ KIỂM TRA", 0)
            draft.add_paragraph(f"Nguồn: {doc['name']} | phiên bản {version} | chế độ {provenance}")
            for label, field in [("Số ký hiệu", value.number), ("Cơ quan", value.agency), ("Ngày văn bản", value.document_date)]:
                draft.add_paragraph(f"{label}: {field.value if field else '[CẦN BỔ SUNG]'}")
            for task in value.tasks:
                draft.add_paragraph(f"Việc đề xuất: {task.request.value}\nHạn nguyên văn: {task.deadline.value if task.deadline else '[CHƯA XÁC ĐỊNH]'}")
            for item in value.missing:
                draft.add_paragraph("[CẦN BỔ SUNG] " + item)
            draft.add_heading("Trích nguồn để kiểm tra", 1)
            draft.add_paragraph(encoded)
            target = self.path("drafts", f"{document_id}-{version}.docx")
            temp = self.path("drafts", f"{document_id}-{version}.tmp")
            draft.save(temp)
            thaw(target)  # Windows: os.replace that bai neu dich dang read-only
            os.replace(temp, target)
            freeze(target)
            draft_hash = hashlib.sha256(target.read_bytes()).hexdigest()
            with self.db() as db:
                db.execute("INSERT INTO versions(document_id,version,content,hash,mode,model,draft_hash) VALUES(?,?,?,?,?,?,?)", (document_id, version, encoded, digest, provenance, self.model if provenance == "ollama" else "", draft_hash))
                db.execute("UPDATE documents SET state='awaiting_review',error='',error_kind='' WHERE id=?", (document_id,))
            return version

    def verify_draft(self, document_id, version, expected_hash):
        """Đối chiếu bytes DOCX trên đĩa với draft_hash đã lưu; trả về đường dẫn.

        Hash rỗng nghĩa là **chưa xác minh**, không phải đã xác minh. Cố ý không
        backfill bằng cách băm tệp đang có trên đĩa: băm sau sự việc chỉ đóng dấu
        lên đúng những byte tình cờ nằm đó, kể cả byte đã bị sửa.
        """
        if not expected_hash:
            raise ValueError("Phiên bản này tạo trước khi có draft_hash nên dự thảo chưa xác minh được; hãy tạo phiên bản mới")
        path = self.path("drafts", f"{document_id}-{version}.docx")
        if not path.is_file():
            raise ValueError("Không tìm thấy tệp dự thảo của phiên bản này; hãy tạo phiên bản mới")
        if hashlib.sha256(path.read_bytes()).hexdigest() != expected_hash:
            raise ValueError("Dự thảo trên đĩa khác bản đã lưu; hãy tạo phiên bản mới")
        return path

    def review(self, document_id, version, digest, action, reason):
        if action not in {"approved", "rejected"}:
            raise ValueError("Hành động không hợp lệ")
        with self.lock, self.db() as db:
            db.execute("BEGIN IMMEDIATE")
            state = db.execute("SELECT state FROM documents WHERE id=?", (document_id,)).fetchone()
            # Kiem su ton tai truoc: khong co tai lieu ma bao "phien ban cu hoac
            # da duyet" thi vua sai ma trang thai vua noi sai su that.
            if state is None:
                raise NotFound("Không tìm thấy tài liệu")
            latest = db.execute("SELECT * FROM versions WHERE document_id=? ORDER BY version DESC LIMIT 1", (document_id,)).fetchone()
            if not latest or latest["version"] != version or latest["hash"] != digest or state["state"] != "awaiting_review":
                raise Conflict("Phiên bản cũ hoặc đã duyệt; tải lại trang")
            if action == "approved":
                # Hash noi dung JSON khong noi gi ve tep DOCX. Nguoi dung xac nhan
                # cai ho doc trong DOCX, nen dung bytes DOCX moi la thu phai khop.
                # Chi chan approve: neu chan ca reject thi mot ban draft bi sua se
                # ket lai o awaiting_review vinh vien, khong the tu choi.
                self.verify_draft(document_id, version, latest["draft_hash"])
            db.execute("INSERT INTO approvals(document_id,version,hash,action,reason) VALUES(?,?,?,?,?)", (document_id, version, digest, action, reason[:2000]))
            db.execute("UPDATE documents SET state=? WHERE id=?", (action, document_id))

    def tasks(self):
        result = []
        for doc in self.listing():
            full = self.get(doc["id"])
            if full["latest"]:
                for task in json.loads(full["latest"]["content"])["tasks"]:
                    result.append({"document_id": doc["id"], "version": full["latest"]["version"], "state": doc["state"], **task})
        return result
