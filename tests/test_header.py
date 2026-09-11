from io import BytesIO
from docx import Document
from tro_ly_van_ban.header import extract_header
from tro_ly_van_ban.parser import parse


def txt(*lines):
    return parse("\n".join(lines).encode(), ".txt")[0]


def values(blocks):
    found = extract_header(blocks)
    for item in found.values():
        if item:
            # Khong bia: gia tri nam trong quote, quote nam trong dung khoi nguon.
            source = next(b["text"] for b in blocks if b["id"] == item["block_id"])
            assert item["value"] in item["quote"] and item["quote"] in source
    return {k: v and v["value"] for k, v in found.items()}


def test_standard_txt_header():
    assert values(txt("UBND HUYỆN GIẢ LẬP A", "Số: 145/UBND-VP", "Giả Lập A, ngày 05 tháng 9 năm 2026",
                      "Kính gửi: Các phòng")) == {
        "number": "145/UBND-VP", "agency": "UBND HUYỆN GIẢ LẬP A", "document_date": "ngày 05 tháng 9 năm 2026"}


def test_national_motto_urgency_and_doc_type_are_not_agency():
    got = values(txt("UBND TỈNH GIẢ LẬP Q", "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", "Độc lập - Tự do - Hạnh phúc",
                     "KHẨN", "Số: 2210/UBND-VX", "Giả Lập Q, ngày 15/09/2026"))
    assert got == {"number": "2210/UBND-VX", "agency": "UBND TỈNH GIẢ LẬP Q", "document_date": "ngày 15/09/2026"}
    assert values(txt("TRƯỜNG TIỂU HỌC GIẢ LẬP C", "THÔNG BÁO", "Giả Lập C, ngày 04 tháng 9 năm 2026"))["agency"] \
        == "TRƯỜNG TIỂU HỌC GIẢ LẬP C"


def test_issuing_body_below_parent_body():
    # ND 30/2020: co quan chu quan o tren, co quan ban hanh o duoi.
    assert values(txt("UBND TỈNH GIẢ LẬP X", "SỞ Y TẾ GIẢ LẬP", "Số: 12/SYT-KHTC"))["agency"] == "SỞ Y TẾ GIẢ LẬP"


def test_number_without_colon_and_with_year_segment():
    assert values(txt("HỘI GIẢ LẬP", "Số 15/KH-HND"))["number"] == "15/KH-HND"
    assert values(txt("CÔNG TY GIẢ LẬP", "Số: 09/2026/CV-GLM"))["number"] == "09/2026/CV-GLM"


def test_body_reference_is_not_taken_as_this_document():
    got = values(txt("SỞ NÔNG NGHIỆP GIẢ LẬP N", "Kính gửi: UBND các huyện",
                     "Thực hiện Công văn số 77/UBND-KT ngày 20/08/2026 của UBND tỉnh, Sở đề nghị báo cáo."))
    assert got["number"] is None and got["document_date"] is None


def test_nothing_matches_returns_none_not_a_guess():
    assert values(txt("Kính gửi anh chị", "Nội dung trao đổi thường ngày.")) == {
        "number": None, "agency": None, "document_date": None}


def test_pdf_page_block_with_merged_two_column_line():
    # PDF: ca trang mot khoi nhieu dong; hai cot bi gop tren mot dong.
    blocks = [{"id": "b1", "location": "Trang 1", "text":
               "UBND HUYỆN GIẢ LẬP A CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\n"
               "Độc lập - Tự do - Hạnh phúc\n"
               "Số: 145/UBND-VP Giả Lập A, ngày 05 tháng 9 năm 2026\n"
               "Kính gửi: Các phòng, ban"}]
    assert values(blocks) == {"number": "145/UBND-VP", "agency": "UBND HUYỆN GIẢ LẬP A",
                              "document_date": "ngày 05 tháng 9 năm 2026"}


def test_docx_header_table_is_read_before_body():
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "BAN QUẢN LÝ GIẢ LẬP E", "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM"
    table.cell(1, 0).text, table.cell(1, 1).text = "Số: 17/GM-BQLDA", "Giả Lập E, ngày 06/09/2026"
    document.add_paragraph("KẾT QUẢ RÀ SOÁT CÔNG TRÌNH")  # dong viet hoa trong than, khong phai co quan
    document.add_paragraph("Đề nghị các phòng rà soát.")
    stream = BytesIO()
    document.save(stream)
    blocks = parse(stream.getvalue(), ".docx")[0]
    assert blocks[0]["location"].startswith("Đoạn")  # parser dat bang sau doan van
    assert values(blocks) == {"number": "17/GM-BQLDA", "agency": "BAN QUẢN LÝ GIẢ LẬP E",
                              "document_date": "ngày 06/09/2026"}
