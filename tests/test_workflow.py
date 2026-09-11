import asyncio
import json
import os
import re
import subprocess
import sys
from io import BytesIO
import httpx
import pytest
from docx import Document
from pypdf import PdfWriter
from fastapi.testclient import TestClient
from fastmcp import Client
from tro_ly_van_ban.domain import Extraction, validate_evidence
from tro_ly_van_ban.fsguard import freeze, thaw
from tro_ly_van_ban.service import Service
from tro_ly_van_ban.model import ModelUnavailable
from tro_ly_van_ban.mcp_server import create_mcp
from tro_ly_van_ban.web import create_app


@pytest.fixture
def service(tmp_path):
    return Service(tmp_path, mode="demo")


def sample(service):
    return service.ingest("mau.txt", "Số: 12/ABC\nGửi báo cáo trước ngày 20/09/2026.".encode())


def content():
    return {"number": {"value": "12/ABC", "block_id": "b1", "quote": "Số: 12/ABC"}, "tasks": [{"request": {"value": "Gửi báo cáo", "block_id": "b2", "quote": "Gửi báo cáo trước ngày 20/09/2026."}, "deadline": {"value": "20/09/2026", "block_id": "b2", "quote": "Gửi báo cáo trước ngày 20/09/2026."}}]}


def test_dedup_persistence_and_version_bound_approval(service):
    doc_id = sample(service)
    assert service.ingest("renamed.txt", "Số: 12/ABC\nGửi báo cáo trước ngày 20/09/2026.".encode()) == doc_id
    assert len(service.listing()) == 1
    service.save(doc_id, content())
    old = service.get(doc_id)["latest"]
    service.save(doc_id, content(), expected_version=1)
    with pytest.raises(ValueError):
        service.review(doc_id, 1, old["hash"], "approved", "")
    latest = service.get(doc_id)["latest"]
    service.review(doc_id, 2, latest["hash"], "approved", "Đã đối chiếu")
    reopened = Service(service.root, mode="demo")
    assert reopened.get(doc_id)["state"] == "approved"
    assert reopened.tasks()[0]["deadline"]["value"] == "20/09/2026"
    assert reopened.get(doc_id)["approvals"][0]["reason"] == "Đã đối chiếu"
    assert Document(service.path("drafts", f"{doc_id}-2.docx")).paragraphs
    with pytest.raises(ValueError):
        service.save(doc_id, content(), expected_version=1)


def test_fabricated_source_rejected(service):
    doc_id = sample(service)
    value = content()
    value["number"]["value"] = "99/FAKE"
    with pytest.raises(ValueError):
        service.save(doc_id, value)
    value = content()
    value["number"]["block_id"] = "../other"
    with pytest.raises(ValueError):
        service.save(doc_id, value)
    assert service.get(doc_id)["latest"] is None


def test_missing_model_persisted(service, monkeypatch):
    service.mode = "ollama"
    def offline(*args, **kwargs):
        raise httpx.ConnectError("offline")
    monkeypatch.setattr(httpx.Client, "post", offline)
    doc_id = sample(service)
    with pytest.raises(ModelUnavailable):
        service.run(doc_id)
    assert Service(service.root).get(doc_id)["state"] == "model_unavailable"


def test_demo_and_scan(service):
    doc_id = sample(service)
    service.run(doc_id)
    value = json.loads(service.get(doc_id)["latest"]["content"])
    assert value["tasks"] == []
    assert "DEMO" in value["missing"][0]
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    data = BytesIO()
    writer.write(data)
    scan = service.ingest("scan.pdf", data.getvalue())
    assert service.get(scan)["state"] == "needs_ocr"
    with pytest.raises(ValueError):
        service.run(scan)


def test_web_csrf_origin_and_review(service):
    with TestClient(create_app(service), base_url="http://127.0.0.1") as client:
        page = client.get("/")
        csrf = re.search('name="csrf" value="([^"]+)"', page.text).group(1)
        assert client.post("/upload", files={"file": ("x.txt", b"test")}).status_code == 403
        assert client.post("/upload", data={"csrf": csrf}, headers={"origin": "https://evil.example"}, files={"file": ("x.txt", b"test")}).status_code == 403
        response = client.post("/upload", data={"csrf": csrf}, files={"file": ("x.txt", b"test")})
        assert response.status_code == 200
        assert "x.txt" in response.text


def test_mcp_client_and_no_approval_tool(service):
    doc_id = sample(service)
    async def run():
        async with Client(create_mcp(service)) as client:
            names = {t.name for t in await client.list_tools()}
            assert names == {"list_documents", "read_document", "get_evidence"}
            result = await client.call_tool("get_evidence", {"document_id": doc_id, "block_ids": ["b1"]})
            assert "12/ABC" in str(result)
    asyncio.run(run())


def test_path_and_file_limits(service, tmp_path):
    with pytest.raises(ValueError):
        service.path("originals", "../../escape")
    with pytest.raises(ValueError):
        service.ingest("old.doc", b"anything")
    with pytest.raises(ValueError):
        service.ingest("huge.txt", b"x" * (10 * 1024 * 1024 + 1))


def test_symlink_in_store_rejected(service, tmp_path):
    outside = tmp_path / "outside"
    outside.write_text("keep")
    link = service.root / "originals" / "link"
    try:
        link.symlink_to(outside)
    except OSError as exc:  # Windows doi quyen SeCreateSymbolicLink
        pytest.skip(f"Không tạo được symlink trong môi trường này: {exc}")
    with pytest.raises(ValueError):
        service.path("originals", "link")


@pytest.mark.skipif(sys.platform != "win32", reason="Junction chỉ có trên Windows")
def test_junction_in_store_rejected(service, tmp_path):
    # Junction khong phai symlink nen is_symlink() bo qua; chan boi kiem tra thu muc cha.
    outside = tmp_path / "outside_dir"
    outside.mkdir()
    link = service.root / "originals" / "junction"
    made = subprocess.run(["cmd", "/c", "mklink", "/J", str(link), str(outside)], capture_output=True)
    if made.returncode != 0:
        pytest.skip("Không tạo được junction: " + made.stderr.decode("utf-8", "replace")[:200])
    with pytest.raises(ValueError):
        service.path("originals", "junction")


def test_frozen_file_can_be_replaced_and_removed(tmp_path):
    # Windows: read-only chan ca os.replace lan os.remove; thaw phai go duoc.
    target = tmp_path / "t.bin"
    target.write_bytes(b"cu")
    freeze(target)
    source = tmp_path / "t.tmp"
    source.write_bytes(b"moi")
    thaw(target)
    os.replace(source, target)
    assert target.read_bytes() == b"moi"
    freeze(target)
    thaw(target)
    target.unlink()
    assert not target.exists()


def test_db_file_not_locked_after_use(service):
    # Windows giu file handle neu connection sqlite khong duoc dong.
    doc_id = sample(service)
    service.run(doc_id)
    service.listing()
    database = service.root / "state.sqlite3"
    moved = service.root / "state.backup"
    os.replace(database, moved)
    assert moved.exists()


def test_edit_form_and_reject_flow(service):
    doc_id = sample(service)
    service.run(doc_id)
    with TestClient(create_app(service), base_url="http://127.0.0.1") as client:
        response = client.get(f"/documents/{doc_id}")
        csrf = re.search('name="csrf" value="([^"]+)"', response.text).group(1)
        response = client.post(f"/documents/{doc_id}/edit", data={
            "csrf": csrf, "version": 1, "number": "12/ABC", "number_block": "b1",
            "request_0": "Gửi báo cáo", "request_0_block": "b2",
            "deadline_0": "20/09/2026", "deadline_0_block": "b2",
        })
        assert response.status_code == 200
        latest = service.get(doc_id)["latest"]
        assert latest["version"] == 2
        response = client.post(f"/documents/{doc_id}/review", data={
            "csrf": csrf, "version": 2, "hash": latest["hash"], "action": "rejected", "reason": "Cần sửa",
        })
        assert service.get(doc_id)["state"] == "rejected"
        assert service.tasks()[0]["state"] == "rejected"
        with pytest.raises(ValueError):
            service.review(doc_id, 2, latest["hash"], "approved", "")


def test_docx_parser_and_bad_input(service):
    document = Document()
    document.add_paragraph("Yêu cầu báo cáo.")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Dữ liệu bảng"
    stream = BytesIO()
    document.save(stream)
    doc_id = service.ingest("test.docx", stream.getvalue())
    blocks = service.get(doc_id)["blocks"]
    assert blocks[0]["location"] == "Đoạn 1"
    assert blocks[1]["location"] == "Bảng 1, hàng 1"
    with pytest.raises(ValueError, match="Không đọc được file"):
        service.ingest("broken.pdf", b"broken")


def test_large_model_input_does_not_call_network(service, monkeypatch):
    service.mode = "ollama"
    doc_id = service.ingest("long.txt", ("dữ liệu " * 1000).encode())
    def forbidden(*args, **kwargs):
        raise AssertionError("Network must not be called")
    monkeypatch.setattr(httpx.Client, "post", forbidden)
    with pytest.raises(ValueError, match="5.000"):
        service.run(doc_id)
    assert service.get(doc_id)["state"] == "error"


def test_review_requires_matching_draft_bytes(service):
    # Hash noi dung JSON khong bao ve tep DOCX. Nguoi dung xac nhan cai ho doc
    # trong DOCX, nen sua DOCX sau khi luu phai chan duoc xac nhan.
    doc_id = sample(service)
    version = service.save(doc_id, content())
    latest = service.get(doc_id)["latest"]
    draft = service.path("drafts", f"{doc_id}-{version}.docx")
    thaw(draft)
    draft.write_bytes(b"DOCX da bi thay the")
    freeze(draft)
    with pytest.raises(ValueError, match="khác bản đã lưu"):
        service.review(doc_id, version, latest["hash"], "approved", "")
    assert service.get(doc_id)["state"] == "awaiting_review"
    assert service.get(doc_id)["approvals"] == []
    # Tu choi van phai chay duoc, neu khong ban draft bi sua se ket lai vinh vien.
    service.review(doc_id, version, latest["hash"], "rejected", "Tệp không khớp")
    assert service.get(doc_id)["state"] == "rejected"


def test_missing_draft_file_blocks_approval(service):
    doc_id = sample(service)
    version = service.save(doc_id, content())
    latest = service.get(doc_id)["latest"]
    draft = service.path("drafts", f"{doc_id}-{version}.docx")
    thaw(draft)
    draft.unlink()
    with pytest.raises(ValueError, match="Không tìm thấy tệp dự thảo"):
        service.review(doc_id, version, latest["hash"], "approved", "")
    assert service.get(doc_id)["state"] == "awaiting_review"


def test_empty_draft_hash_is_unverified_not_verified(service):
    # Migration dat draft_hash='' cho ban cu. Rong = chua xac minh, khong duoc
    # coi la da xac minh, va khong duoc backfill bang cach bam tep dang co.
    doc_id = sample(service)
    version = service.save(doc_id, content())
    latest = service.get(doc_id)["latest"]
    with service.db() as db:
        db.execute("UPDATE versions SET draft_hash='' WHERE document_id=? AND version=?", (doc_id, version))
    assert service.get(doc_id)["latest"]["draft_hash"] == ""
    with pytest.raises(ValueError, match="chưa xác minh"):
        service.review(doc_id, version, latest["hash"], "approved", "")
    assert service.get(doc_id)["state"] == "awaiting_review"
    with TestClient(create_app(service), base_url="http://127.0.0.1") as client:
        page = client.get(f"/documents/{doc_id}/draft/{version}")
        assert "chưa xác minh được" in page.text
        assert "không xác minh được" in client.get(f"/documents/{doc_id}").text


def test_needs_ocr_blocked_in_service_not_only_ui(service):
    # Cac route van nhan POST truc tiep du giao dien da an nut, nen hang rao
    # phai nam o Service.save.
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    data = BytesIO()
    writer.write(data)
    scan = service.ingest("scan.pdf", data.getvalue())
    assert service.get(scan)["state"] == "needs_ocr"
    for call in (lambda: service.run(scan),
                 lambda: service.save(scan, {}),
                 lambda: service.save(scan, {}, provenance="manual")):
        with pytest.raises(ValueError, match="OCR"):
            call()
    assert service.get(scan)["state"] == "needs_ocr"
    assert service.get(scan)["latest"] is None
    with TestClient(create_app(service), base_url="http://127.0.0.1") as client:
        csrf = re.search('name="csrf" value="([^"]+)"', client.get("/").text).group(1)
        for route, payload in (("manual", {"version": 0}),
                               ("save", {"version": 0, "content": "{}"}),
                               ("edit", {"version": 0})):
            response = client.post(f"/documents/{scan}/{route}", data={"csrf": csrf, **payload})
            assert "OCR" in response.text, route
    assert service.get(scan)["state"] == "needs_ocr"
    assert service.get(scan)["latest"] is None
